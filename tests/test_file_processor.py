"""
Automated unit tests for file processing and document extraction.
"""

import io
from pathlib import Path
import pytest
from werkzeug.datastructures import FileStorage
from file_processor import (
    allowed_file,
    extract_text_from_txt,
    process_uploaded_file,
    FileProcessingError
)
import docx
import pypdf


def test_allowed_file():
    """Verify that only supported file extensions are permitted."""
    assert allowed_file("lecture_notes.pdf") is True
    assert allowed_file("assignment.docx") is True
    assert allowed_file("summary.txt") is True
    assert allowed_file("notes.md") is True
    assert allowed_file("script.py") is False
    assert allowed_file("malicious.exe") is False
    assert allowed_file("archive.zip") is False
    assert allowed_file("") is False
    assert allowed_file("no_extension") is False


def test_extract_text_from_txt(tmp_path):
    """Verify plain text and markdown extraction."""
    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("Binary search requires a sorted array with O(log n) complexity.", encoding="utf-8")

    result = extract_text_from_txt(txt_file)
    assert result["success"] is True
    assert "Binary search" in result["text"]
    assert result["word_count"] > 0
    assert result["character_count"] > 0


def test_empty_txt_file(tmp_path):
    """Verify empty text file returns appropriate error."""
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("", encoding="utf-8")

    result = extract_text_from_txt(empty_file)
    assert result["success"] is False
    assert "empty" in result["error"].lower()


def test_process_docx_file(tmp_path):
    """Verify Word DOCX file extraction."""
    docx_path = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_heading("Operating Systems Review", level=1)
    doc.add_paragraph("A semaphore is a synchronization tool.")
    doc.save(str(docx_path))

    with open(docx_path, "rb") as f:
        file_storage = FileStorage(
            stream=io.BytesIO(f.read()),
            filename="test.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        result = process_uploaded_file(file_storage, tmp_path / "temp_uploads")

    assert result["success"] is True
    assert "Operating Systems" in result["text"]
    assert "semaphore" in result["text"]


def test_process_pdf_file(tmp_path):
    """Verify PDF file extraction using pypdf."""
    pdf_path = tmp_path / "test.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with open(pdf_path, "wb") as f:
        writer.write(f)

    with open(pdf_path, "rb") as f:
        file_storage = FileStorage(
            stream=io.BytesIO(f.read()),
            filename="test.pdf",
            content_type="application/pdf"
        )
        result = process_uploaded_file(file_storage, tmp_path / "temp_uploads")

    # A blank PDF page has < 30 characters and triggers the scanned/empty OCR warning
    assert result["success"] is False
    assert result["is_scanned"] is True
    assert "ocr" in result["error"].lower()


def test_unsupported_file_upload(tmp_path):
    """Verify unsupported file types are rejected."""
    file_storage = FileStorage(
        stream=io.BytesIO(b"binary content"),
        filename="payload.exe",
        content_type="application/octet-stream"
    )

    with pytest.raises(FileProcessingError) as exc_info:
        process_uploaded_file(file_storage, tmp_path / "temp_uploads")

    assert "unsupported" in str(exc_info.value).lower()
